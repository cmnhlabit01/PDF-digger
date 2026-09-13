import io
import os
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import pymupdf as fitz
from PIL import Image

from src.config import DEFAULT_FONT, FALLBACK_MODELS
from src.docx_builder import DocxBuilder
from src.gemini_extractor import GeminiExtractor, clean_thai_ocr_text
from src.pdf_processor import ExtractedImage, PDFProcessor, enhance_document_image


class TestPDFDigger(unittest.TestCase):
    def setUp(self):
        self.test_dir = Path("./test_output")
        self.test_dir.mkdir(exist_ok=True)

    def tearDown(self):
        # ลบไฟล์ทดสอบ
        for f in self.test_dir.glob("*"):
            try:
                f.unlink()
            except Exception:
                pass
        try:
            self.test_dir.rmdir()
        except Exception:
            pass

    def test_docx_builder_thai_font_and_tables(self):
        """ทดสอบการสร้างไฟล์ Word: ตรวจสอบฟอนต์ Cordia New, โครงสร้างตาราง และการแทรกรูป"""
        builder = DocxBuilder(font_name="Cordia New")

        markdown_content = (
            "# รายงานสรุปผลการดำเนินงานประจำปี\n"
            "## ข้อมูลทั่วไปและการทดสอบภาษาไทย\n"
            "นี่คือข้อความภาษาไทยที่มีสระบน สระล่าง เช่น ที่นี่ ปี่ น้้ำ และ **ข้อความตัวหนา** กับ *ตัวเอียง*\n\n"
            "| ลำดับ | รายการ | จำนวน | ราคาต่อหน่วย | รวมเป็นเงิน |\n"
            "|---|---|---|---|---|\n"
            "| 1 | ค่าบริการจัดทำเอกสาร | 2 | 1,500.00 | 3,000.00 |\n"
            "| 2 | ค่าแปลภาษาไทย-อังกฤษ | 5 | 800.00 | 4,000.00 |\n\n"
            "[IMAGE]\n\n"
            "- รายการหัวข้อย่อยที่ 1\n"
            "- รายการหัวข้อย่อยที่ 2\n"
        )

        # สร้างภาพจำลอง
        img_buffer = io.BytesIO()
        pil_img = Image.new("RGB", (200, 100), color="blue")
        pil_img.save(img_buffer, format="PNG")
        dummy_image = ExtractedImage(
            page_num=0,
            image_index=1,
            image_bytes=img_buffer.getvalue(),
            ext="png",
            width=200,
            height=100,
        )

        builder.add_page_content(
            markdown_text=markdown_content,
            page_num=1,
            images=[dummy_image],
            is_first_page=True,
        )

        out_path = self.test_dir / "test_doc.docx"
        builder.save(out_path)

        self.assertTrue(out_path.exists())

        # ตรวจสอบโครงสร้างไฟล์ docx ที่สร้างขึ้น
        saved_doc = docx.Document(str(out_path))

        # ตรวจสอบว่ามีตารางอย่างน้อย 1 ตาราง
        self.assertEqual(len(saved_doc.tables), 1)
        tbl = saved_doc.tables[0]
        self.assertEqual(len(tbl.rows), 3)  # หัวตาราง 1 แถว + ข้อมูล 2 แถว
        self.assertEqual(len(tbl.columns), 5)

        # ตรวจสอบหัวตาราง
        self.assertEqual(tbl.rows[0].cells[1].text.strip(), "รายการ")
        self.assertEqual(tbl.rows[1].cells[1].text.strip(), "ค่าบริการจัดทำเอกสาร")

        # ตรวจสอบว่ามีการตั้งค่า Cordia New ใน XML Complex Script (cs)
        xml_content = saved_doc._body._element.xml
        self.assertIn('w:cs="Cordia New"', xml_content)
        self.assertIn('w:ascii="Cordia New"', xml_content)
        print("\n✅ ทดสอบ DocxBuilder (ฟอนต์ Cordia New, ตาราง, รูปภาพ) สำเร็จ")

    def test_pdf_processor_rendering(self):
        """ทดสอบการเปิดและเรนเดอร์หน้า PDF เป็นภาพ 200 DPI"""
        test_pdf = self.test_dir / "sample.pdf"

        # สร้าง PDF ตัวอย่างด้วย PyMuPDF
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((50, 72), "Test Document / ทดสอบเอกสารภาษาไทย", fontsize=16)

        # แทรกรูปภาพทดสอบลงใน PDF
        img_buffer = io.BytesIO()
        pil_img = Image.new("RGB", (100, 100), color="red")
        pil_img.save(img_buffer, format="PNG")
        page.insert_image(fitz.Rect(50, 100, 150, 200), stream=img_buffer.getvalue())

        doc.save(str(test_pdf))
        doc.close()

        processor = PDFProcessor(test_pdf, dpi=200)
        self.assertEqual(processor.get_page_count(), 1)

        page_data = processor.process_page(0)
        self.assertGreater(len(page_data.rendered_image_bytes), 0)
        self.assertGreaterEqual(len(page_data.embedded_images), 1)
        self.assertEqual(page_data.embedded_images[0].ext, "png")
        print("✅ ทดสอบ PDFProcessor (เรนเดอร์ 200 DPI และสกัดภาพ) สำเร็จ")

    def test_gemini_fallback_mechanism(self):
        """ทดสอบระบบ Seamless Fallback เมื่อโมเดลแรกติด 429 Quota Exceeded"""
        fallback_log = []

        def fallback_cb(prev, nxt, reason):
            fallback_log.append((prev, nxt, reason))

        with patch("src.gemini_extractor.genai.Client") as mock_client_cls:
            mock_client = MagicMock()
            mock_client_cls.return_value = mock_client

            call_count = 0

            def mock_generate_content(model, contents, config):
                nonlocal call_count
                call_count += 1
                if model == "gemini-3.6-flash":
                    # จำลองข้อผิดพลาด HTTP 429 Quota Exceeded
                    raise Exception("429 ResourceExhausted: Quota exceeded for model gemini-3.6-flash")
                else:
                    # โมเดลสำรองทำงานสำเร็จ
                    mock_resp = MagicMock()
                    mock_resp.text = "# หัวข้อที่แกะได้จากโมเดลสำรอง\nเนื้อหาทดสอบ"
                    return mock_resp

            mock_client.models.generate_content.side_effect = mock_generate_content

            extractor = GeminiExtractor(
                api_key="test-api-key",
                models=["gemini-3.6-flash", "gemini-3.7-flash", "gemini-flash-latest"],
                on_fallback=fallback_cb,
            )

            # ทดสอบเรียกประมวลผล
            dummy_bytes = b"fake-image-png-bytes"
            result_text, used_model, _ = extractor.extract_page_markdown(dummy_bytes, 0)

            # ตรวจสอบว่าสลับไปใช้ gemini-3.7-flash สำเร็จ
            self.assertEqual(used_model, "gemini-3.7-flash")
            self.assertIn("หัวข้อที่แกะได้จากโมเดลสำรอง", result_text)

            # ตรวจสอบว่ามีการบันทึกเหตุการณ์ fallback
            self.assertEqual(len(fallback_log), 1)
            self.assertEqual(fallback_log[0][0], "gemini-3.6-flash")
            self.assertEqual(fallback_log[0][1], "gemini-3.7-flash")
            self.assertIn("429 Quota Exceeded", fallback_log[0][2])

            print("✅ ทดสอบ Seamless Model Fallback (สลับโมเดลอัตโนมัติเมื่อโควตาเต็ม) สำเร็จ")

    def test_font_detector_normalization(self):
        """ทดสอบการทำความสะอาดและตรวจจับชื่อฟอนต์มาตรฐาน"""
        from src.font_detector import clean_font_name, detect_dominant_font

        # 1. ทดสอบการตัด prefix และแปลงชื่อฟอนต์
        self.assertEqual(clean_font_name("ABCDEF+THSarabunPSK-Bold"), "TH Sarabun New")
        self.assertEqual(clean_font_name("CordiaNew-Regular"), "Cordia New")
        self.assertEqual(clean_font_name("AngsanaUPC,Bold"), "Angsana New")
        self.assertEqual(clean_font_name("BrowalliaNew"), "Browallia New")
        self.assertEqual(clean_font_name("Tahoma-Bold"), "Tahoma")
        self.assertEqual(clean_font_name("Calibri"), "Calibri")

        # 2. ทดสอบตรวจจับฟอนต์จาก PDF ดิจิทัล
        test_pdf = self.test_dir / "font_sample.pdf"
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((50, 72), "This is a document with Helvetica or Times font", fontname="helv", fontsize=14)
        doc.save(str(test_pdf))
        doc.close()

        detected = detect_dominant_font(test_pdf)
        self.assertIsNotNone(detected)
        self.assertEqual(detected, "Arial")  # helv mapped to Arial
        print("✅ ทดสอบ Font Detector (ตรวจจับและแปลงชื่อฟอนต์ต้นฉบับ) สำเร็จ")

    def test_image_support_and_enhancement(self):
        """ทดสอบการรองรับไฟล์รูปภาพโดยตรง (.png) และระบบ Image Enhancement"""
        test_img_path = self.test_dir / "sample_scan.png"
        img = Image.new("RGB", (800, 600), color=(240, 235, 220))  # กระดาษอมเหลือง
        img.save(test_img_path)

        # 1. ตรวจสอบว่า PDFProcessor จัดการรูปภาพได้ถูกต้อง
        processor = PDFProcessor(test_img_path)
        self.assertTrue(processor.is_image)
        self.assertEqual(processor.get_page_count(), 1)
        self.assertIsNone(processor.detect_font())

        # 2. ตรวจสอบการแปลงและการเปิดโหมด enhance
        p_data = processor.process_page(0, enhance=True)
        self.assertGreater(len(p_data.rendered_image_bytes), 0)

        # 3. ตรวจสอบฟังก์ชัน enhance_document_image โดยตรง
        enhanced_bytes = enhance_document_image(p_data.rendered_image_bytes)
        self.assertGreater(len(enhanced_bytes), 0)
        print("✅ ทดสอบ Direct Image Input & Image Enhancement สำเร็จ")

    def test_translation_prompt_generation(self):
        """ทดสอบการสร้างคำสั่งแปลภาษา (Target Language) ใน GeminiExtractor"""
        with patch("src.gemini_extractor.genai.Client") as mock_client_cls:
            mock_client = MagicMock()
            mock_client_cls.return_value = mock_client

            captured_prompt = None

            def mock_generate_content(model, contents, config):
                nonlocal captured_prompt
                captured_prompt = contents[1]  # prompt string
                mock_resp = MagicMock()
                mock_resp.text = "# Translated Content in Thai\nเนื้อหาที่แปลเป็นไทย"
                return mock_resp

            mock_client.models.generate_content.side_effect = mock_generate_content

            extractor = GeminiExtractor(api_key="test-api-key")
            dummy_bytes = b"fake-image-bytes"

            # ทดสอบแปลเป็นไทย
            result, _, _ = extractor.extract_page_markdown(dummy_bytes, 0, target_language="th")
            self.assertIn("แปลเนื้อหาทั้งหมดเป็นภาษาไทย", captured_prompt)
            self.assertIn("Translated Content", result)

            # ทดสอบแปลเป็นอังกฤษ
            result_en, _, _ = extractor.extract_page_markdown(dummy_bytes, 0, target_language="en")
            self.assertIn("แปลเนื้อหาทั้งหมดเป็นภาษาอังกฤษ", captured_prompt)

            print("✅ ทดสอบ Translation Prompt Generation (แปลไทย/อังกฤษ) สำเร็จ")

    def test_layout_alignment_and_image_bbox(self):
        """ทดสอบการจัดตำแหน่ง (Center, Right, Justify), ขนาดตัวอักษร Small/Metric Scaling และการจัดตำแหน่งรูปภาพจาก bbox"""
        builder = DocxBuilder(font_name="Cordia New")

        # 1. ทดสอบการถอดรหัสแท็กจัดหน้า (parse_line_formatting)
        raw_center = "[CENTER]# หัวข้อกึ่งกลางหน้ากระดาษ[/CENTER]"
        cleaned, align, is_small = builder.parse_line_formatting(raw_center)
        self.assertEqual(cleaned, "# หัวข้อกึ่งกลางหน้ากระดาษ")
        self.assertEqual(align, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertFalse(is_small)

        raw_right_small = "[RIGHT][SMALL]วันที่ ๑๕ มกราคม ๒๕๖๗[/SMALL][/RIGHT]"
        cleaned, align, is_small = builder.parse_line_formatting(raw_right_small)
        self.assertEqual(cleaned, "วันที่ ๑๕ มกราคม ๒๕๖๗")
        self.assertEqual(align, WD_ALIGN_PARAGRAPH.RIGHT)
        self.assertTrue(is_small)

        # 2. ทดสอบ Metric Font Scaling เมื่อสลับไปใช้ Calibri
        builder_calibri = DocxBuilder(font_name="Calibri")
        self.assertEqual(builder_calibri.font_sizes["body"], 11.5)
        self.assertEqual(builder_calibri.font_sizes["h1"], 16.0)
        self.assertEqual(builder.font_sizes["body"], 16.0)
        self.assertEqual(builder.font_sizes["h1"], 22.0)

        # 3. ทดสอบการสร้างพารากราฟและการวางตำแหน่งรูปภาพตาม Bounding Box (1:1 Sizing)
        markdown_content = (
            "[CENTER]# บันทึกข้อความ[/CENTER]\n"
            "[RIGHT]ส่วนราชการ สำนักนายกรัฐมนตรี[/RIGHT]\n"
            "[JUSTIFY]ด้วยสำนักงานมีความประสงค์จะจัดประชุมสัมมนาเชิงปฏิบัติการเพื่อพัฒนาระบบเทคโนโลยีสารสนเทศ[/JUSTIFY]\n"
            "[SMALL]หมายเหตุ: ข้อความขนาดเล็กส่วนท้ายเอกสาร[/SMALL]\n"
            "[IMAGE]\n"
            "[IMAGE]\n"
        )

        # จำลองรูปภาพ 2 รูป: รูปแรกอยู่ตรงกลาง (ตราครุฑ), รูปที่สองอยู่มุมขวา (ตรายางรับหนังสือ)
        img_buffer = io.BytesIO()
        pil_img = Image.new("RGB", (100, 100), color="green")
        pil_img.save(img_buffer, format="PNG")
        raw_img_bytes = img_buffer.getvalue()

        # รูปที่ 1: x0=250, x1=345 (ตรงกลางหน้า A4 กว้าง ~595 pt)
        img_center = ExtractedImage(
            page_num=0,
            image_index=1,
            image_bytes=raw_img_bytes,
            ext="png",
            width=95,
            height=95,
            bbox=(250.0, 50.0, 345.0, 145.0),
        )

        # รูปที่ 2: x0=450, x1=530 (มุมขวาบน)
        img_right = ExtractedImage(
            page_num=0,
            image_index=2,
            image_bytes=raw_img_bytes,
            ext="png",
            width=80,
            height=80,
            bbox=(450.0, 50.0, 530.0, 130.0),
        )

        builder.add_page_content(
            markdown_text=markdown_content,
            page_num=1,
            images=[img_center, img_right],
            is_first_page=True,
        )

        out_path = self.test_dir / "test_layout_doc.docx"
        builder.save(out_path)
        self.assertTrue(out_path.exists())

        doc_check = docx.Document(str(out_path))
        # ตรวจสอบการจัดตำแหน่งพารากราฟ
        # Paragraph 0: Center Heading
        self.assertEqual(doc_check.paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertEqual(doc_check.paragraphs[0].text.strip(), "บันทึกข้อความ")

        # Paragraph 1: Right aligned metadata
        self.assertEqual(doc_check.paragraphs[1].alignment, WD_ALIGN_PARAGRAPH.RIGHT)
        self.assertEqual(doc_check.paragraphs[1].text.strip(), "ส่วนราชการ สำนักนายกรัฐมนตรี")

        # Paragraph 2: Justified body paragraph
        self.assertEqual(doc_check.paragraphs[2].alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)

        # Paragraph 3: Small text size check (13pt for Cordia)
        self.assertAlmostEqual(doc_check.paragraphs[3].runs[0].font.size.pt, 13.0)

        # ตรวจสอบการจัดตำแหน่งของรูปภาพ:
        # Paragraph 4: Image Center
        self.assertEqual(doc_check.paragraphs[4].alignment, WD_ALIGN_PARAGRAPH.CENTER)

        # Paragraph 5: Image Right
        self.assertEqual(doc_check.paragraphs[5].alignment, WD_ALIGN_PARAGRAPH.RIGHT)

        print("✅ ทดสอบ Layout Alignment, Font Scaling & Image 1:1 Bbox Sizing สำเร็จ")

    def test_table_cell_formatting_and_image(self):
        """ทดสอบการประมวลผลภายในเซลล์ตาราง: แยกบรรทัด <br>, แท็ก [CENTER], ตัวหนา และการแทรกรูปภาพในเซลล์"""
        builder = DocxBuilder(font_name="Cordia New")

        # ตรวจสอบว่าระยะขอบหน้ากระดาษเป็น 0.5 นิ้ว
        for section in builder.doc.sections:
            self.assertAlmostEqual(section.top_margin.inches, 0.5)
            self.assertAlmostEqual(section.left_margin.inches, 0.5)

        markdown_table = (
            "| ข้อมูลผู้ส่งและผู้รับ | รายละเอียดบาร์โค้ด |\n"
            "|---|---|\n"
            "| **ผู้ส่ง (FROM)** **กฤต**<br>2 ถนน สุเทพ (ชั้น 3)<br>**ผู้รับ (TO)** **Huang** | [IMAGE]<br>[CENTER]**TH2608793615733**[/CENTER] |\n"
        )

        img_buffer = io.BytesIO()
        pil_img = Image.new("RGB", (120, 60), color="black")
        pil_img.save(img_buffer, format="PNG")
        dummy_barcode = ExtractedImage(
            page_num=0,
            image_index=1,
            image_bytes=img_buffer.getvalue(),
            ext="png",
            width=120,
            height=60,
            bbox=(350.0, 40.0, 520.0, 90.0),
        )

        builder.add_page_content(
            markdown_text=markdown_table,
            page_num=1,
            images=[dummy_barcode],
            is_first_page=True,
        )

        out_path = self.test_dir / "test_table_cell_doc.docx"
        builder.save(out_path)
        self.assertTrue(out_path.exists())

        saved_doc = docx.Document(str(out_path))
        self.assertEqual(len(saved_doc.tables), 1)
        tbl = saved_doc.tables[0]

        # ตรวจสอบเซลล์ฝั่งซ้าย (แถวที่ 1 คอลัมน์ที่ 0): ต้องถูกแยกเป็น 3 paragraphs จาก <br>
        left_cell = tbl.rows[1].cells[0]
        self.assertGreaterEqual(len(left_cell.paragraphs), 3)
        self.assertNotIn("<br>", left_cell.text)
        self.assertNotIn("**", left_cell.text)
        self.assertIn("ผู้ส่ง (FROM)", left_cell.paragraphs[0].text)

        # ตรวจสอบเซลล์ฝั่งขวา (แถวที่ 1 คอลัมน์ที่ 1): ต้องมีรูปภาพบาร์โค้ด และข้อความ TH... ที่จัดกึ่งกลางโดยไม่มีแท็กดิบหลุด
        right_cell = tbl.rows[1].cells[1]
        self.assertNotIn("[CENTER]", right_cell.text)
        self.assertNotIn("[/CENTER]", right_cell.text)
        self.assertNotIn("**", right_cell.text)
        self.assertIn("TH2608793615733", right_cell.text)

        # ตรวจสอบว่าบรรทัดรหัสถูกจัดกึ่งกลาง
        code_p = [p for p in right_cell.paragraphs if "TH2608793615733" in p.text][0]
        self.assertEqual(code_p.alignment, WD_ALIGN_PARAGRAPH.CENTER)

        # ทดสอบการสร้างตารางไร้ขอบ [BORDERLESS]
        builder_bl = DocxBuilder()
        builder_bl.add_page_content(
            "[BORDERLESS]\n"
            "| [IMAGE] | **TH2608793615733** |\n"
            "|---|---|\n",
            page_num=1,
            is_first_page=True,
        )
        self.assertEqual(len(builder_bl.doc.tables), 1)
        bl_tbl = builder_bl.doc.tables[0]
        # ตรวจสอบว่ามี element w:tblBorders อยู่ใน tblPr
        tblPr = bl_tbl._tbl.tblPr
        tblBorders = tblPr.find(qn("w:tblBorders"))
        self.assertIsNotNone(tblBorders)

        print("✅ ทดสอบ Table Cell Formatting (<br>, [CENTER], Bold, [IMAGE], [BORDERLESS]) สำเร็จ")

    def test_thai_ocr_cleaner_and_image_sorting(self):
        """ทดสอบฟังก์ชันแก้คำผิดวรรณยุกต์ไทย และการจัดเรียงรูปภาพตามพิกัดสายตา (y0, x0)"""
        # 1. ทดสอบการทำความสะอาดวรรณยุกต์ไทยและคำผิดในแบบฟอร์ม
        raw_typos = (
            "กฤต อยู่ชั2น 3 วันที2 11 ขั#นตอนการส่งคืน เพืLอความรวดเร็ว ช้อปปี2 "
            "เพิ2มเติม สิ3นสุด และไม่ตอ้งเกบ็ เงิน พันกงานผู้รับสินค้า จังหัวดเชียงใหม่ "
            "หัลกฐาน ขั(นตอน นีE สิEนสุด ชื@อ"
        )
        cleaned = clean_thai_ocr_text(raw_typos)
        self.assertIn("ชั้น 3", cleaned)
        self.assertIn("วันที่ 11", cleaned)
        self.assertIn("ขั้นตอนการส่งคืน", cleaned)
        self.assertIn("เพื่อความรวดเร็ว", cleaned)
        self.assertIn("ช้อปปี้", cleaned)
        self.assertIn("เพิ่มเติม", cleaned)
        self.assertIn("สิ้นสุด", cleaned)
        self.assertIn("ไม่ต้องเก็บ เงิน", cleaned)
        self.assertIn("พนักงานผู้รับสินค้า", cleaned)
        self.assertIn("จังหวัดเชียงใหม่", cleaned)
        self.assertIn("หลักฐาน", cleaned)
        self.assertIn("ขั้นตอน", cleaned)
        self.assertIn("นี้", cleaned)
        self.assertIn("ชื่อ", cleaned)

        # 2. ทดสอบการจัดเรียงรูปภาพแบบ Row Band Clustering (y0 ใกล้เคียงกัน ต้องเรียงจากซ้ายไปขวา x0)
        test_pdf = self.test_dir / "test_sort_images.pdf"
        doc = fitz.open()
        page = doc.new_page(width=595, height=842)

        # สร้างภาพ 3 รูป:
        # รูปขวาบน (y=6.7, x=273.5) เช่น บาร์โค้ด
        # รูปซ้ายบน (y=9.7, x=18.7) เช่น โลโก้
        # รูปล่าง (y=500, x=50)
        img_buf = io.BytesIO()
        Image.new("RGB", (50, 50), color="blue").save(img_buf, format="PNG")
        page.insert_image(fitz.Rect(273.5, 6.7, 500.0, 50.0), stream=img_buf.getvalue())  # ขวาบน (y เล็กกว่าเล็กน้อย)
        page.insert_image(fitz.Rect(18.7, 9.7, 150.0, 50.0), stream=img_buf.getvalue())   # ซ้ายบน (y มากกว่าเล็กน้อยแต่อยู่แถวเดียวกัน)
        page.insert_image(fitz.Rect(50.0, 500.0, 100.0, 550.0), stream=img_buf.getvalue()) # รูปล่าง
        doc.save(str(test_pdf))
        doc.close()

        proc = PDFProcessor(test_pdf)
        pdata = proc.process_page(0)

        # ตรวจสอบว่ารูปภาพในแถวบน โลโก้ซ้าย (x=18.7) ต้องมาก่อน บาร์โค้ดขวา (x=273.5) แม้ y บาร์โค้ดจะเริ่มก่อนเล็กน้อย
        self.assertEqual(len(pdata.embedded_images), 3)
        img_top_left = pdata.embedded_images[0]
        img_top_right = pdata.embedded_images[1]
        img_bottom = pdata.embedded_images[2]

        self.assertLess(img_top_left.bbox[0], img_top_right.bbox[0])
        self.assertAlmostEqual(img_top_left.bbox[0], 18.7, delta=1.0)
        self.assertAlmostEqual(img_top_right.bbox[0], 273.5, delta=1.0)
        self.assertGreater(img_bottom.bbox[1], 400.0)

        self.assertEqual(img_top_left.image_index, 1)
        self.assertEqual(img_top_right.image_index, 2)
        self.assertEqual(img_bottom.image_index, 3)

        # 3. ทดสอบการตัดทอนเส้นใต้ลายเซ็นที่ยาวเกินไปใน DocxBuilder
        builder = DocxBuilder()
        builder.add_page_content(
            "ชื่อผู้ส่ง: __________________________________________________\n"
            "✂---------------------------------------------------------------------------------------------------",
            page_num=1,
            is_first_page=True,
        )
        p_sig = builder.doc.paragraphs[0]
        self.assertNotIn("__________________________________________________", p_sig.text)
        self.assertIn("____________________", p_sig.text)

        p_cut = builder.doc.paragraphs[1]
        self.assertLess(len(p_cut.text), 70)

        print("✅ ทดสอบ Thai OCR Cleaner, Row Band Image Sorting & Signature Line Trimming สำเร็จ")

    def test_table_merging_badges_and_headings(self):
        """ทดสอบการผสานเซลล์ตารางแนวตั้ง ป้ายกำกับ [BADGE] หัวข้อในตาราง และการแยกตารางที่อยู่ติดกัน"""
        builder = DocxBuilder(font_name="Cordia New")

        md_content = (
            "| [BADGE]PICK UP[/BADGE] | [CENTER]**SPX**[/CENTER] |\n"
            "|---|---|\n\n"
            "| **ผู้รับ (TO)** | # **I17-(PET.5)** |\n"
            "|---|---|\n"
            "| | [CENTER]**-**[/CENTER] |\n"
            "| | [CENTER]**RR-MP**[/CENTER] |\n"
            "| **บาร์โค้ด** | [BADGE]W[/BADGE] |\n"
            "| | [CENTER]# **10**[/CENTER] |\n\n"
            "Original Order No: 260906BRSPB0KA &nbsp;&nbsp;&nbsp;&nbsp; Pickup Date: 11-09-2026\n"
        )

        builder.add_page_content(
            markdown_text=md_content,
            page_num=1,
            is_first_page=True,
        )

        out_path = self.test_dir / "test_badges_tables.docx"
        builder.save(out_path)
        self.assertTrue(out_path.exists())

        doc = docx.Document(str(out_path))
        # 1. ตรวจสอบว่าตารางทั้งสองไม่ถูกหลอมรวมกัน (มี 2 ตารางแยกกัน)
        self.assertEqual(len(doc.tables), 2)

        # 2. ตรวจสอบป้ายกำกับ PICK UP ในตารางที่ 1 (มีพื้นหลังทึบ 595959 และสีตัวอักษรขาว FFFFFF)
        tbl1 = doc.tables[0]
        cell_pickup = tbl1.rows[0].cells[0]
        self.assertIn('w:fill="595959"', cell_pickup._tc.xml)
        self.assertIn('w:color w:val="FFFFFF"', cell_pickup._tc.xml)
        self.assertIn("PICK UP", cell_pickup.text)

        # 3. ตรวจสอบการผสานเซลล์แนวตั้งในตารางที่ 2: ผู้รับ (TO) ควบแถว 0, 1, 2
        tbl2 = doc.tables[1]
        # เซลล์แถว 0 และแถว 1 ในคอลัมน์ 0 ต้องเป็นเซลล์เดียวกันหลังการ merge
        self.assertEqual(tbl2.rows[0].cells[0].text.strip(), tbl2.rows[1].cells[0].text.strip())

        # 4. ตรวจสอบหัวข้อขนาดใหญ่ในเซลล์ (# I17-(PET.5)) ต้องมีขนาดอย่างน้อย 20pt
        cell_i17 = tbl2.rows[0].cells[1]
        run_i17 = [r for r in cell_i17.paragraphs[0].runs if r.text][0]
        self.assertGreaterEqual(run_i17.font.size.pt, 20.0)

        # 5. ตรวจสอบป้ายกำกับ W ในตารางที่ 2 (แถวที่ 3)
        cell_w = tbl2.rows[3].cells[1]
        self.assertIn('w:fill="595959"', cell_w._tc.xml)
        self.assertIn('w:color w:val="FFFFFF"', cell_w._tc.xml)

        # 6. ตรวจสอบข้อความ receipt ว่าไม่มี &nbsp; หลุดออกมา
        p_receipt = [p for p in doc.paragraphs if "Original Order No" in p.text][0]
        self.assertNotIn("&nbsp;", p_receipt.text)
        self.assertIn("Original Order No: 260906BRSPB0KA", p_receipt.text)
        self.assertIn("Pickup Date: 11-09-2026", p_receipt.text)

        print("✅ ทดสอบ Table Vertical Merging, Badges, Cell Headings & Entity Decoding สำเร็จ")

    def test_uneven_table_columns_handling(self):
        """ทดสอบการจัดการตาราง Markdown ที่แต่ละแถวมีจำนวนคอลัมน์ไม่เท่ากัน เพื่อป้องกัน IndexError: list index out of range"""
        builder = DocxBuilder(font_name="Cordia New")

        # ตารางที่แถว 1 มี 4 คอลัมน์, แถว 2 มี 2 คอลัมน์, แถว 3 มี 1 คอลัมน์ (เช่น แถวสรุปผล)
        md_content = (
            "| รหัส | สินค้า | จำนวน | ราคา |\n"
            "|---|---|---|---|\n"
            "| 01 | ดินสอ | 2 | 20 |\n"
            "| หมายเหตุ: ชำระผ่านบัตร |\n"
            "| รวมสุทธิ | 20 |\n"
        )

        # ต้องทำงานได้โดยไม่ raise IndexError
        builder.add_page_content(
            markdown_text=md_content,
            page_num=1,
            is_first_page=True,
        )

        out_path = self.test_dir / "test_uneven_table.docx"
        builder.save(out_path)
        self.assertTrue(out_path.exists())

        doc = docx.Document(str(out_path))
        self.assertEqual(len(doc.tables), 1)
        tbl = doc.tables[0]
        # ต้องมี 4 คอลัมน์เท่ากับจำนวนคอลัมน์สูงสุด
        self.assertEqual(len(tbl.columns), 4)
        print("✅ ทดสอบ Uneven Table Columns Handling (ป้องกัน list index out of range) สำเร็จ")


if __name__ == "__main__":
    unittest.main()
